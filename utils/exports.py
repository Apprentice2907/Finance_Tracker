import datetime
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from db.transactions import get_transactions, get_totals, get_category_totals, bulk_insert_transactions
from db.categories import get_categories, get_or_create_category

def export_to_excel(file_path: str, start_date=None, end_date=None, period_label="All Time") -> str:
    """
    Exports transactions, summary statistics, and category breakdown into a styled Excel workbook.
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    header_fill = PatternFill(start_color="2962D6", end_color="2962D6", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    sub_fill = PatternFill(start_color="F0F4FD", end_color="F0F4FD", fill_type="solid")
    sub_font = Font(name="Segoe UI", size=11, bold=True, color="18212F")
    regular_font = Font(name="Segoe UI", size=10, color="18212F")
    green_font = Font(name="Segoe UI", size=10, bold=True, color="159B72")
    red_font = Font(name="Segoe UI", size=10, bold=True, color="D65B67")
    thin_side = Side(border_style="thin", color="E7EAF0")
    border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    # ==========================
    # SHEET 1: Transactions
    # ==========================
    ws_trans = wb.create_sheet(title="Transactions")
    ws_trans.views.sheetView[0].showGridLines = True
    ws_trans.freeze_panes = "A2"

    trans_headers = ["ID", "Date", "Name", "Type", "Category", "Amount (₹)", "Note"]
    ws_trans.append(trans_headers)

    for col_idx, _ in enumerate(trans_headers, 1):
        cell = ws_trans.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_idx in (1, 2, 4) else ("right" if col_idx == 6 else "left"), vertical="center")

    rows = get_transactions(start_date, end_date, sort_by="date_desc")
    for row_idx, t in enumerate(rows, 2):
        t_id, t_type, t_amount, _, t_cat, t_date, t_note, _, *rest = t
        t_name = rest[0] if rest else ""
        ws_trans.append([t_id, t_date, t_name or "", t_type.title(), t_cat, t_amount, t_note or ""])
        for c_idx in range(1, 8):
            c = ws_trans.cell(row=row_idx, column=c_idx)
            c.border = border_all
            c.font = regular_font
            if c_idx in (1, 2, 4):
                c.alignment = Alignment(horizontal="center")
                if c_idx == 4:
                    c.font = green_font if t_type == "income" else red_font
            elif c_idx == 6:
                c.number_format = '"₹"#,##0.00'
                c.alignment = Alignment(horizontal="right")
                c.font = green_font if t_type == "income" else red_font

    ws_trans.auto_filter.ref = f"A1:G{max(2, len(rows) + 1)}"

    # Auto-adjust column widths
    for col in ws_trans.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_trans.column_dimensions[col_letter].width = max(max_len + 4, 12)
    ws_trans.column_dimensions["E"].width = 16
    ws_trans.column_dimensions["F"].width = 28

    # ==========================
    # SHEET 2: Summary
    # ==========================
    ws_sum = wb.create_sheet(title="Summary")
    ws_sum.views.sheetView[0].showGridLines = True
    
    ws_sum.append(["Financial Summary Report"])
    ws_sum.merge_cells("A1:B1")
    ws_sum.cell(row=1, column=1).font = Font(name="Segoe UI", size=14, bold=True, color="2962D6")
    
    totals = get_totals(start_date, end_date)
    income = totals.get("income", 0.0) or 0.0
    expense = totals.get("expense", 0.0) or 0.0
    net = income - expense

    summary_data = [
        ("Report Period", period_label),
        ("Date Range", f"{start_date or 'Start'} to {end_date or 'Present'}"),
        ("Generated On", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Total Income", income),
        ("Total Expenses", expense),
        ("Net Balance", net),
        ("Transaction Count", len(rows)),
    ]

    ws_sum.append([]) # Blank row
    for row_idx, (label, val) in enumerate(summary_data, 3):
        ws_sum.append([label, val])
        c1 = ws_sum.cell(row=row_idx, column=1)
        c2 = ws_sum.cell(row=row_idx, column=2)
        c1.font = sub_font
        c1.fill = sub_fill
        c1.border = border_all
        c2.font = regular_font
        c2.border = border_all
        if isinstance(val, (int, float)) and label.startswith("Total") or label.startswith("Net"):
            c2.number_format = '"₹"#,##0.00'
            c2.font = green_font if (val >= 0 and label != "Total Expenses") else red_font
            c2.alignment = Alignment(horizontal="right")

    ws_sum.column_dimensions["A"].width = 24
    ws_sum.column_dimensions["B"].width = 28

    # ==========================
    # SHEET 3: Categories
    # ==========================
    ws_cat = wb.create_sheet(title="Categories")
    ws_cat.views.sheetView[0].showGridLines = True
    ws_cat.freeze_panes = "A2"

    cat_headers = ["Category", "Type", "Total Amount (₹)", "Transaction Count"]
    ws_cat.append(cat_headers)
    for col_idx, _ in enumerate(cat_headers, 1):
        cell = ws_cat.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_idx <= 2 else "right", vertical="center")

    expense_cats = get_category_totals("expense", start_date, end_date)
    income_cats = get_category_totals("income", start_date, end_date)
    
    current_cat_row = 2
    for name, amt, _, cnt in income_cats:
        ws_cat.append([name, "Income", amt, cnt])
        for c_idx in range(1, 5):
            c = ws_cat.cell(row=current_cat_row, column=c_idx)
            c.border = border_all
            c.font = regular_font
            if c_idx == 2:
                c.font = green_font
                c.alignment = Alignment(horizontal="center")
            elif c_idx == 3:
                c.number_format = '"₹"#,##0.00'
                c.alignment = Alignment(horizontal="right")
                c.font = green_font
            elif c_idx == 4:
                c.alignment = Alignment(horizontal="right")
        current_cat_row += 1

    for name, amt, _, cnt in expense_cats:
        ws_cat.append([name, "Expense", amt, cnt])
        for c_idx in range(1, 5):
            c = ws_cat.cell(row=current_cat_row, column=c_idx)
            c.border = border_all
            c.font = regular_font
            if c_idx == 2:
                c.font = red_font
                c.alignment = Alignment(horizontal="center")
            elif c_idx == 3:
                c.number_format = '"₹"#,##0.00'
                c.alignment = Alignment(horizontal="right")
                c.font = red_font
            elif c_idx == 4:
                c.alignment = Alignment(horizontal="right")
        current_cat_row += 1

    for col in ws_cat.columns:
        col_letter = get_column_letter(col[0].column)
        ws_cat.column_dimensions[col_letter].width = 22

    wb.save(file_path)
    return file_path


def parse_and_import_excel(file_path: str, skip_duplicates: bool = True) -> dict:
    """
    Parses an Excel (.xlsx) file and imports valid transaction records.
    Returns:
    {
        "success": True/False,
        "total_rows": int,
        "imported": int,
        "skipped_duplicates": int,
        "errors": list of (row_num, error_msg),
        "message": str
    }
    """
    result = {
        "success": False,
        "total_rows": 0,
        "imported": 0,
        "skipped_duplicates": 0,
        "errors": [],
        "message": ""
    }

    if not os.path.exists(file_path):
        result["message"] = "Selected file does not exist."
        return result

    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        # Choose Transactions sheet if available, else first sheet
        sheet_name = "Transactions" if "Transactions" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet_name]

        # Find header row
        header_row = None
        date_col = None
        type_col = None
        cat_col = None
        amount_col = None
        name_col = None
        note_col = None

        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if not row or not any(row):
                continue
            cleaned_cells = [str(c).strip().lower() if c is not None else "" for c in row]
            
            # Check if this row looks like header
            for col_idx, cell_val in enumerate(cleaned_cells):
                if "date" in cell_val:
                    date_col = col_idx
                elif "name" in cell_val or "title" in cell_val or "item" in cell_val or "merchant" in cell_val:
                    name_col = col_idx
                elif "type" in cell_val or "kind" in cell_val:
                    type_col = col_idx
                elif "category" in cell_val or "cat" in cell_val:
                    cat_col = col_idx
                elif "amount" in cell_val or "value" in cell_val or "price" in cell_val:
                    amount_col = col_idx
                elif "note" in cell_val or "desc" in cell_val or "memo" in cell_val:
                    note_col = col_idx
            
            if date_col is not None and amount_col is not None:
                header_row = row_idx
                break

        if header_row is None:
            result["message"] = "Could not find valid column headers (at least Date and Amount required)."
            return result

        transactions_to_insert = []
        
        for row_idx, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            if not row or not any(row):
                continue
            
            result["total_rows"] += 1
            
            # 1. Parse Date
            raw_date = row[date_col] if date_col is not None and date_col < len(row) else None
            parsed_date = None
            if isinstance(raw_date, (datetime.date, datetime.datetime)):
                parsed_date = raw_date.strftime("%Y-%m-%d")
            elif isinstance(raw_date, str) and raw_date.strip():
                raw_str = raw_date.strip()
                # Try common date formats
                for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d", "%m/%d/%Y", "%d %b %Y", "%d %B %Y"):
                    try:
                        parsed_date = datetime.datetime.strptime(raw_str, fmt).strftime("%Y-%m-%d")
                        break
                    except ValueError:
                        continue
            
            if not parsed_date:
                result["errors"].append((row_idx, f"Invalid date format: '{raw_date}'"))
                continue

            # 2. Parse Amount
            raw_amount = row[amount_col] if amount_col is not None and amount_col < len(row) else None
            parsed_amount = None
            if isinstance(raw_amount, (int, float)):
                parsed_amount = float(abs(raw_amount))
            elif isinstance(raw_amount, str):
                # Clean symbols like ₹, $, commas
                clean_str = raw_amount.replace("₹", "").replace("$", "").replace(",", "").strip()
                try:
                    parsed_amount = float(abs(float(clean_str)))
                except ValueError:
                    pass

            if parsed_amount is None or parsed_amount <= 0:
                result["errors"].append((row_idx, f"Invalid amount: '{raw_amount}'"))
                continue

            # 3. Parse Type
            raw_type = row[type_col] if type_col is not None and type_col < len(row) else "expense"
            parsed_type = "expense"
            if raw_type:
                s_type = str(raw_type).strip().lower()
                if "inc" in s_type or "credit" in s_type or "in" == s_type:
                    parsed_type = "income"
                else:
                    parsed_type = "expense"

            # 4. Parse Category
            raw_cat = row[cat_col] if cat_col is not None and cat_col < len(row) else None
            cat_id = None
            if raw_cat and str(raw_cat).strip() and str(raw_cat).strip() != "—" and str(raw_cat).strip().lower() != "uncategorised":
                cat_name = str(raw_cat).strip()
                cat_id = get_or_create_category(cat_name, parsed_type)

            # 5. Parse Note with Formula Injection Protection
            raw_note = row[note_col] if note_col is not None and note_col < len(row) else ""
            parsed_note = str(raw_note).strip() if raw_note is not None else ""
            if parsed_note and parsed_note[0] in ("=", "+", "-", "@", "\t", "\r"):
                parsed_note = "'" + parsed_note

            # 6. Parse Name
            raw_name = row[name_col] if name_col is not None and name_col < len(row) else ""
            parsed_name = str(raw_name).strip() if raw_name is not None else ""
            if parsed_name and parsed_name[0] in ("=", "+", "-", "@", "\t", "\r"):
                parsed_name = "'" + parsed_name

            transactions_to_insert.append((parsed_type, parsed_amount, cat_id, parsed_date, parsed_note, parsed_name))

    except Exception as e:
        result["message"] = f"Failed to read Excel workbook: {str(e)}"
        return result
    finally:
        try:
            wb.close()
        except Exception:
            pass

    # Bulk insert with transaction safety
    if transactions_to_insert:
        ins, skp = bulk_insert_transactions(transactions_to_insert, skip_duplicates=skip_duplicates)
        result["imported"] = ins
        result["skipped_duplicates"] = skp
        result["invalid_rows"] = len(result["errors"])
        result["success"] = True
        result["message"] = f"Imported: {ins} | Skipped duplicates: {skp} | Invalid rows: {len(result['errors'])} | Total processed: {result['total_rows']}"
        
        # Record in observability registry
        from utils.observability import METRICS
        METRICS.import_stats["imported"] += ins
        METRICS.import_stats["skipped_duplicates"] += skp
        METRICS.import_stats["invalid"] += len(result["errors"])
    else:
        result["message"] = f"No valid transaction rows found to import (Total rows evaluated: {result['total_rows']}, Errors: {len(result['errors'])})."

    return result

def export_to_csv(file_path: str, start_date=None, end_date=None) -> str:
    """
    Exports all transactions into a standard UTF-8 CSV format with formula injection protection.
    """
    import csv
    from db.transactions import get_transactions

    rows = get_transactions(start_date, end_date, sort_by="date_desc")
    with open(file_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["ID", "Date", "Name", "Type", "Category", "Amount", "Note"])
        for t in rows:
            t_id, t_type, t_amount, _, t_cat, t_date, t_note, _, *rest = t
            t_name = rest[0] if rest else ""
            safe_name = t_name or ""
            if safe_name and safe_name[0] in ("=", "+", "-", "@", "\t", "\r"):
                safe_name = "'" + safe_name
            safe_note = t_note or ""
            if safe_note and safe_note[0] in ("=", "+", "-", "@", "\t", "\r"):
                safe_note = "'" + safe_note
            writer.writerow([t_id, t_date, safe_name, t_type.title(), t_cat, f"{t_amount:.2f}", safe_note])
    
    from utils.observability import METRICS
    METRICS.export_stats["exported_rows"] += len(rows)
    METRICS.export_stats["export_count"] += 1
    return file_path
