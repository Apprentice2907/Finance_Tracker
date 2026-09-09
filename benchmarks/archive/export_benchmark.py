import time
import os
import sys
import tracemalloc
import sqlite3
import openpyxl
import json
from typing import Dict, Any

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from db.database import get_connection

def run_export_benchmark(db_path: str, row_limit: int = 10_000) -> Dict[str, Any]:
    print(f"\n--- Benchmarking Excel Export on {row_limit:,} records ---")
    conn = sqlite3.connect(db_path)
    cur = conn.cursor()
    cur.execute("SELECT id, type, amount, category_id, date, note FROM transactions LIMIT ?", (row_limit,))
    rows = cur.fetchall()
    conn.close()

    out_std = f"benchmarks/results/export_std_{row_limit}.xlsx"
    out_stream = f"benchmarks/results/export_stream_{row_limit}.xlsx"

    # 1. Standard openpyxl Workbook
    tracemalloc.start()
    t0 = time.perf_counter()
    wb_std = openpyxl.Workbook()
    ws_std = wb_std.active
    ws_std.title = "Transactions"
    ws_std.append(["ID", "Type", "Amount", "Category ID", "Date", "Note"])
    for r in rows:
        ws_std.append(list(r))
    wb_std.save(out_std)
    t1 = time.perf_counter()
    _, peak_mem_std = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    std_time = t1 - t0
    std_size = os.path.getsize(out_std) / 1024.0 # KB

    # 2. Streaming write_only openpyxl Workbook
    tracemalloc.start()
    t0 = time.perf_counter()
    wb_stream = openpyxl.Workbook(write_only=True)
    ws_stream = wb_stream.create_sheet(title="Transactions")
    ws_stream.append(["ID", "Type", "Amount", "Category ID", "Date", "Note"])
    for r in rows:
        ws_stream.append(list(r))
    wb_stream.save(out_stream)
    t1 = time.perf_counter()
    _, peak_mem_stream = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    stream_time = t1 - t0
    stream_size = os.path.getsize(out_stream) / 1024.0 # KB

    if os.path.exists(out_std): os.remove(out_std)
    if os.path.exists(out_stream): os.remove(out_stream)

    res = {
        "rows": len(rows),
        "standard_export": {
            "time_sec": round(std_time, 3),
            "throughput_rows_sec": round(len(rows) / std_time, 1),
            "peak_ram_mb": round(peak_mem_std / (1024 * 1024), 2),
            "file_size_kb": round(std_size, 1)
        },
        "streaming_export": {
            "time_sec": round(stream_time, 3),
            "throughput_rows_sec": round(len(rows) / stream_time, 1),
            "peak_ram_mb": round(peak_mem_stream / (1024 * 1024), 2),
            "file_size_kb": round(stream_size, 1)
        }
    }

    print(f"  Standard Openpyxl: {std_time:.3f}s ({res['standard_export']['throughput_rows_sec']:,.0f} rows/s) | Peak RAM: {res['standard_export']['peak_ram_mb']} MB")
    print(f"  Streaming (write_only): {stream_time:.3f}s ({res['streaming_export']['throughput_rows_sec']:,.0f} rows/s) | Peak RAM: {res['streaming_export']['peak_ram_mb']} MB")
    return res

if __name__ == "__main__":
    db_100k = os.path.join(os.path.dirname(__file__), "datasets", "transactions_100k.db")
    res_10k = run_export_benchmark(db_100k, row_limit=10_000)
    res_50k = run_export_benchmark(db_100k, row_limit=50_000)
    with open("benchmarks/results/export_benchmark_results.json", "w") as f:
        json.dump({"10k": res_10k, "50k": res_50k}, f, indent=2)
